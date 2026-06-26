"""Spreadsheet ingest + column auto-detect, ported from the donor's index.html
(SPEC §7, §15). Server-side here (openpyxl) instead of the donor's in-browser
DecompressionStream reader, but the field patterns, the EU-decimal-comma coord
parsing, the "skip rows without id/lat/lng", and "never invent coordinates"
rules are carried verbatim.
"""
from __future__ import annotations

import csv
import io
import re
from typing import Any

from openpyxl import load_workbook

# --- field definitions (ported SCREEN_FIELDS / SHOP_FIELDS) -----------------
# Each: (key, [regex patterns], required). First matching header wins.
SCREEN_FIELDS = [
    ("screenId", [r"^frame\s*id$", r"^location\s*id$", r"^screen[\s_]*id$", r"^panel\s*id$", r"^site\s*number$", r"(^|[\s_])id$"], True),
    ("lat", [r"^lat", r"latitude"], True),
    ("lng", [r"^lon", r"^lng", r"longitude"], True),
    ("name", [r"location\s*name", r"site\s*name", r"^name$", r"frame\s*name"], False),
    ("city", [r"^city\s*de$", r"^city$", r"^city\s*en$", r"^stadt$", r"^ort$"], False),
    ("plz", [r"^postcode$", r"^plz$", r"^zip", r"postal"], False),
    ("address", [r"address\s*/\s*street", r"^address$", r"^street", r"stra[sß]"], False),
    ("format", [r"frame\s*size", r"venue\s*type", r"^channel$", r"^category$", r"^environment$", r"format"], False),
    ("partnerCol", [r"^publisher$", r"^environment$", r"operator", r"^partner$"], False),
]

# Custom-POI upload (hybrid input, SPEC §4.3) — adapted from donor SHOP_FIELDS.
POI_FIELDS = [
    ("id", [r"^poi[\s_]*id$", r"^shop[\s_]*row[\s_]*id$", r"^id$", r"(^|[\s_])id$"], True),
    ("lat", [r"^latitude2$", r"^lat", r"latitude"], True),
    ("lng", [r"^longitude3$", r"^lon", r"^lng", r"longitude"], True),
    ("name", [r"^name$", r"^bezeichnung$", r"name"], False),
    ("city", [r"^ort$", r"^city", r"^stadt$"], False),
    ("plz", [r"^plz$", r"^postcode$", r"^zip"], False),
    ("category", [r"^category$", r"^kategorie$"], False),
]


def auto_detect(headers: list[str], fields: list[tuple]) -> dict[str, str]:
    """First header matching each field's patterns (case-insensitive)."""
    out: dict[str, str] = {}
    for key, pats, _req in fields:
        found = ""
        for pat in pats:
            rx = re.compile(pat, re.IGNORECASE)
            for h in headers:
                if h and rx.search(h):
                    found = h
                    break
            if found:
                break
        out[key] = found
    return out


def parse_coord(v: Any) -> float | None:
    """Tolerant coordinate parse; handles EU decimal comma (donor parseCoord)."""
    if v is None:
        return None
    s = str(v).strip()
    if not s:
        return None
    if "," in s and "." not in s:
        s = s.replace(",", ".")
    try:
        f = float(s)
    except ValueError:
        return None
    return f if -90 <= f <= 180 else (f if f == f else None)  # NaN guard


def _sv(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def read_spreadsheet(data: bytes, filename: str) -> dict:
    """Return {headers, rows, sheet}. Accepts .xlsx (ZIP) or .csv/.txt.

    For .xlsx the worksheet with the most rows is used (donor heuristic).
    """
    name = (filename or "").lower()
    is_zip = len(data) > 1 and data[0] == 0x50 and data[1] == 0x4B  # 'PK'
    if name.endswith((".csv", ".txt")) or not is_zip:
        return _read_csv(data)
    return _read_xlsx(data)


def _read_csv(data: bytes) -> dict:
    text = data.decode("utf-8-sig", errors="replace")
    first = text.split("\n", 1)[0]
    delim = ";" if first.count(";") > first.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    rows = [r for r in reader]
    if not rows:
        return {"headers": [], "rows": [], "sheet": "CSV"}
    headers = [(h or "").strip() for h in rows[0]]
    data_rows = []
    for r in rows[1:]:
        if not any((c or "").strip() for c in r):
            continue
        data_rows.append({headers[j]: (r[j] if j < len(r) else None) for j in range(len(headers))})
    return {"headers": headers, "rows": data_rows, "sheet": "CSV"}


def _read_xlsx(data: bytes) -> dict:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = max(wb.worksheets, key=lambda s: s.max_row or 0)
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return {"headers": [], "rows": [], "sheet": ws.title}
    headers = [(str(h).strip() if h is not None else "") for h in header_row]
    data_rows = []
    for r in rows_iter:
        if not any(c is not None and str(c).strip() != "" for c in r):
            continue
        data_rows.append({headers[j]: (r[j] if j < len(r) else None) for j in range(len(headers))})
    wb.close()
    return {"headers": [h for h in headers if h], "rows": data_rows, "sheet": ws.title}


# --- record builders (ported buildRecords / buildShopRecords) ---------------
def build_screen_records(rows: list[dict], mapping: dict, partner: str) -> dict:
    """Canonical screen records keyed (partner, screen_id); last row wins on dup.
    Rows lacking screenId/lat/lng are skipped (never invented)."""
    def g(row, field):
        col = mapping.get(field)
        return _sv(row.get(col)) if col else None

    seen: dict[tuple, dict] = {}
    skipped = 0
    for row in rows:
        sid = g(row, "screenId")
        lat = parse_coord(row.get(mapping.get("lat"))) if mapping.get("lat") else None
        lng = parse_coord(row.get(mapping.get("lng"))) if mapping.get("lng") else None
        if not sid or lat is None or lng is None:
            skipped += 1
            continue
        name = g(row, "name") or g(row, "address") or f"Screen {sid}"
        seen[(partner, sid)] = {
            "partner": partner, "screen_id": sid, "name": name, "lat": lat, "lng": lng,
            "city": g(row, "city"), "plz": g(row, "plz"),
            "address": g(row, "address"), "format": g(row, "format"),
        }
    recs = list(seen.values())
    return {"recs": recs, "skipped": skipped, "dupes": len(rows) - len(recs) - skipped}


def build_poi_records(rows: list[dict], mapping: dict, list_id: str) -> dict:
    """Custom POIs; id required, coords optional (kept None = un-plottable)."""
    def g(row, field):
        col = mapping.get(field)
        return _sv(row.get(col)) if col else None

    seen: dict[str, dict] = {}
    order: list[str] = []
    no_id = no_coord = 0
    for row in rows:
        raw_id = g(row, "id")
        if not raw_id:
            no_id += 1
            continue
        lat = parse_coord(row.get(mapping.get("lat"))) if mapping.get("lat") else None
        lng = parse_coord(row.get(mapping.get("lng"))) if mapping.get("lng") else None
        if lat is None or lng is None:
            lat = lng = None
            no_coord += 1
        if raw_id not in seen:
            order.append(raw_id)
        seen[raw_id] = {
            "id": f"custom/{list_id}/{raw_id}", "list_id": list_id,
            "name": g(row, "name") or f"POI {raw_id}",
            "category": g(row, "category") or "custom", "subcategory": None,
            "city": g(row, "city"), "plz": g(row, "plz"),
            "lat": lat, "lng": lng,
        }
    recs = [seen[k] for k in order]
    return {"recs": recs, "no_id": no_id, "no_coord": no_coord,
            "dupes": len(rows) - len(recs) - no_id}
